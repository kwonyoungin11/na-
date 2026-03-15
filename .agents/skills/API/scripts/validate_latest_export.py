#!/usr/bin/env python3
"""최신 API 수집 export CSV를 마스터 XLSX와 비교 검증."""
from __future__ import annotations

import argparse
import csv
import glob
import os
import sys
from datetime import datetime

from openpyxl import load_workbook

from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[4]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from nara_workspace_paths import app_dir, master_path

REQUIRED_18 = [
    "공고번호",
    "공고명",
    "발주처",
    "지역제한",
    "순위",
    "업체명",
    "개찰일시",
    "투찰일시",
    "투찰금액",
    "예가대비투찰률",
    "기초대비투찰률",
    "기초대비사정률",
    "예정가격",
    "낙찰하한가",
    "기초금액",
    "사정률",
    "발주처투찰률",
    "낙찰하한가차이",
]


def normalize_header(value: str) -> str:
    return (value or "").replace("\ufeff", "").strip()


def read_latest_export(export_dir: str):
    pattern = os.path.join(export_dir, "나라장터_전처리완료_원천전체_*.csv")
    files = sorted(glob.glob(pattern), key=os.path.getmtime, reverse=True)
    if not files:
        raise FileNotFoundError(f"원천전체 export 파일이 없습니다: {pattern}")
    latest = files[0]

    with open(latest, "r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        header = [normalize_header(x) for x in next(reader)]
        idx = header.index("공고번호") if "공고번호" in header else 0
        notice_set = set()
        rows = 0
        for row in reader:
            rows += 1
            if len(row) > idx and row[idx]:
                notice_set.add(str(row[idx]).strip())

    return {
        "path": latest,
        "mtime": datetime.fromtimestamp(os.path.getmtime(latest)),
        "header": header,
        "rows": rows,
        "notice_set": notice_set,
    }


def read_master(master_path: str):
    workbook = load_workbook(master_path, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header = [normalize_header(x) for x in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    idx = header.index("공고번호") if "공고번호" in header else 0
    notice_set = set()
    rows = 0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        rows += 1
        if row and len(row) > idx and row[idx] is not None:
            notice_set.add(str(row[idx]).strip())
    workbook.close()
    return {
        "path": master_path,
        "header": header,
        "rows": rows,
        "notice_set": notice_set,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--master", default=str(master_path()))
    parser.add_argument("--export-dir", default=str(app_dir() / "exports"))
    args = parser.parse_args()

    exp = read_latest_export(args.export_dir)
    master = read_master(args.master)

    exp_header_set = set(exp["header"])
    master_header_set = set(master["header"])
    req_missing_export = [c for c in REQUIRED_18 if c not in exp_header_set]
    req_missing_master = [c for c in REQUIRED_18 if c not in master_header_set]

    print("=== NARA API PIPELINE VALIDATION ===")
    print("latest_export:", exp["path"])
    print("latest_export_mtime:", exp["mtime"])
    print("export_rows:", exp["rows"])
    print("export_notice_unique:", len(exp["notice_set"]))
    print("master_path:", master["path"])
    print("master_rows:", master["rows"])
    print("master_notice_unique:", len(master["notice_set"]))
    print("required18_missing_in_export:", req_missing_export)
    print("required18_missing_in_master:", req_missing_master)
    print("header_export_minus_master:", [h for h in exp["header"] if h not in master_header_set])
    print("header_master_minus_export:", [h for h in master["header"] if h not in exp_header_set])
    print("notice_in_export_not_in_master:", len(exp["notice_set"] - master["notice_set"]))
    print("notice_in_master_not_in_export:", len(master["notice_set"] - exp["notice_set"]))


if __name__ == "__main__":
    main()
