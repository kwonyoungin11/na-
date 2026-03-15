#!/usr/bin/env python3
"""낙찰 파이프라인 상태 점검 스크립트.

- DB 중복/규칙위반 점검
- 시드 엑셀(전체_통합_데이터.xlsx)과 공고번호 교차비교
"""

from __future__ import annotations

import argparse
import json
import os
import sqlite3
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional

from openpyxl import load_workbook

import sys

PROJECT_ROOT = Path(__file__).resolve().parents[4]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from nara_workspace_paths import app_dir, master_path


def normalize_token(value: object) -> str:
    if value is None:
        return ""
    return "".join(str(value).split()).upper()


def parse_optional_datetime(text: Optional[str]) -> Optional[datetime]:
    if not text:
        return None
    text = str(text).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def find_default_db_path() -> Optional[Path]:
    candidate = app_dir() / "nara-dashboard.db"
    return candidate if candidate.exists() else None


@dataclass
class DbRow:
    notice_no: str
    notice_number: str
    normalized_company_name: str
    company_name: str
    rank: int
    bid_amount: int
    region: str
    open_date: Optional[datetime]


def load_db_rows(db_path: Path, from_dt: Optional[datetime], to_dt: Optional[datetime]) -> list[DbRow]:
    con = sqlite3.connect(str(db_path))
    con.row_factory = sqlite3.Row
    try:
        cur = con.cursor()
        cur.execute(
            """
            SELECT notice_no, notice_number, normalized_company_name, company_name,
                   rank, IFNULL(bid_amount, 0) AS bid_amount, IFNULL(region, '') AS region,
                   open_date
            FROM participants
            """
        )
        rows: list[DbRow] = []
        for item in cur.fetchall():
            open_dt = parse_optional_datetime(item["open_date"])
            if from_dt and open_dt and open_dt < from_dt:
                continue
            if to_dt and open_dt and open_dt > to_dt:
                continue

            rows.append(
                DbRow(
                    notice_no=item["notice_no"] or "",
                    notice_number=item["notice_number"] or "",
                    normalized_company_name=item["normalized_company_name"] or "",
                    company_name=item["company_name"] or "",
                    rank=int(item["rank"] or 0),
                    bid_amount=int(item["bid_amount"] or 0),
                    region=(item["region"] or "").strip(),
                    open_date=open_dt,
                )
            )
        return rows
    finally:
        con.close()


def db_key(row: DbRow) -> str:
    company = row.normalized_company_name or row.company_name
    return "|".join(
        [
            normalize_token(row.notice_number),
            normalize_token(company),
            str(row.rank),
            str(row.bid_amount),
        ]
    )


def load_xlsx_award_notices(xlsx_path: Path) -> set[str]:
    wb = load_workbook(filename=str(xlsx_path), read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        col_map = {str(v).strip(): idx for idx, v in enumerate(header) if v is not None}

        notice_idx = col_map.get("공고번호")
        rank_idx = col_map.get("순위")
        if notice_idx is None or rank_idx is None:
            return set()

        notices: set[str] = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                continue
            notice_no = row[notice_idx]
            rank_val = row[rank_idx]
            if not notice_no:
                continue
            try:
                rank_num = int(round(float(rank_val)))
            except Exception:
                continue
            if rank_num == 1:
                notices.add(normalize_token(notice_no))
        return notices
    finally:
        wb.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="낙찰 파이프라인 품질 점검")
    parser.add_argument("--db", default="", help="SQLite DB 경로")
    parser.add_argument("--xlsx", default=str(master_path()), help="시드 엑셀 경로")
    parser.add_argument("--from", dest="from_dt", default="", help="시작일시(YYYY-MM-DD[ HH:MM[:SS]])")
    parser.add_argument("--to", dest="to_dt", default="", help="종료일시(YYYY-MM-DD[ HH:MM[:SS]])")
    parser.add_argument("--regions", default="전국,대구,경북", help="허용 지역 목록(콤마구분)")
    args = parser.parse_args()

    from_dt = parse_optional_datetime(args.from_dt)
    to_dt = parse_optional_datetime(args.to_dt)

    db_path = Path(args.db).expanduser() if args.db else find_default_db_path()
    if not db_path or not db_path.exists():
        print("[ERROR] DB 파일을 찾지 못했습니다.")
        return 2

    xlsx_path = Path(args.xlsx).expanduser()
    if not xlsx_path.exists():
        print(f"[WARN] 엑셀 파일이 없습니다: {xlsx_path}")

    allowed_regions = {r.strip() for r in args.regions.split(",") if r.strip()}
    db_rows = load_db_rows(db_path, from_dt, to_dt)

    key_counter = Counter(db_key(r) for r in db_rows)
    duplicate_rows = sum(v - 1 for v in key_counter.values() if v > 1)
    invalid_region = sum(1 for r in db_rows if r.region not in allowed_regions)
    non_award_rank = sum(1 for r in db_rows if r.rank != 1)

    db_notice_set = {normalize_token(r.notice_number or r.notice_no) for r in db_rows if (r.notice_number or r.notice_no)}
    xlsx_notice_set = load_xlsx_award_notices(xlsx_path) if xlsx_path.exists() else set()

    missing_in_db = sorted(xlsx_notice_set - db_notice_set)
    extra_in_db = sorted(db_notice_set - xlsx_notice_set) if xlsx_notice_set else []

    result = {
        "db_path": str(db_path),
        "xlsx_path": str(xlsx_path),
        "rows": len(db_rows),
        "duplicate_rows": duplicate_rows,
        "invalid_region_rows": invalid_region,
        "non_award_rank_rows": non_award_rank,
        "xlsx_award_notice_count": len(xlsx_notice_set),
        "db_notice_count": len(db_notice_set),
        "missing_notice_in_db_count": len(missing_in_db),
        "extra_notice_in_db_count": len(extra_in_db),
        "missing_notice_in_db_sample": missing_in_db[:20],
        "extra_notice_in_db_sample": extra_in_db[:20],
    }

    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
