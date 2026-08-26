"""Replace notice rows in sqlite/master with C++ igunsul jsonl records."""
from __future__ import annotations

import json
import sqlite3
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


def load_jsonl(path: Path) -> list[dict]:
    rows = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.strip():
            rows.append(json.loads(line))
    return rows


def notice_keys(notice: str) -> tuple[str, str]:
    base = notice.split("-")[0]
    return notice, base


def insert_participant(cur: sqlite3.Cursor, fx: dict, now: str) -> None:
    cur.execute(
        """
        INSERT INTO participants (
            notice_no, notice_number, notice_name, ordering_agency, region, rank,
            company_name, normalized_company_name, open_date, bid_date,
            bid_amount, bid_rate_vs_estimate, bid_rate_vs_base, adjust_rate,
            estimate_price, lower_limit_price, base_amount, price_adjustment_rate,
            lower_rate, lower_diff, biz_no, ceo_name, remark, is_awarded, updated_at
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
        """,
        (
            fx["notice_no"],
            fx["notice_number"],
            fx["notice_name"],
            fx["ordering_agency"],
            fx["region"],
            fx["rank"],
            fx["company_name"],
            fx["normalized_company_name"],
            fx["open_date"].replace("/", "-"),
            fx.get("bid_date") or None,
            fx.get("bid_amount"),
            fx.get("bid_rate_vs_estimate"),
            fx.get("bid_rate_vs_base"),
            fx.get("adjust_rate"),
            fx.get("estimate_price"),
            fx.get("lower_limit_price"),
            fx.get("base_amount"),
            fx.get("price_adjustment_rate"),
            fx.get("lower_rate"),
            None,
            fx.get("biz_no") or "",
            "",
            "",
            1 if fx.get("is_awarded") else 0,
            now,
        ),
    )


def merge_sqlite(db_path: Path, rows: list[dict]) -> None:
    by_notice: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        by_notice[row["notice_number"]].append(row)

    con = sqlite3.connect(str(db_path))
    cur = con.cursor()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for notice, fixture_rows in by_notice.items():
        full, base = notice_keys(notice)
        rank1 = next(r for r in fixture_rows if r.get("rank") == 1)
        cur.execute(
            """
            DELETE FROM participants
            WHERE notice_number = ? OR notice_number = ? OR notice_no = ?
            """,
            (full, base, base),
        )
        for fx in fixture_rows:
            insert_participant(cur, fx, now)

        cur.execute(
            "DELETE FROM collector_notices WHERE notice_number IN (?, ?) AND notice_type='award'",
            (full, base),
        )
        cur.execute(
            """
            INSERT INTO collector_notices (
                notice_type, notice_number, notice_name, ordering_agency, notice_date,
                region_text, participant_count_text, updated_at
            ) VALUES ('award', ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                full,
                rank1["notice_name"],
                rank1["ordering_agency"],
                rank1["open_date"].replace("/", "-"),
                rank1["region"],
                f"{len(fixture_rows)}개사",
                now,
            ),
        )

    con.commit()
    con.close()


def master_row(headers: list, fx: dict, participant_count: int) -> list:
    mapping = {
        "공고번호": fx["notice_number"],
        "공사명": fx["notice_name"],
        "발주처": fx["ordering_agency"],
        "지역제한": fx["region"],
        "개찰일시": fx["open_date"].replace("/", "-"),
        "기초금액": fx.get("base_amount"),
        "투찰률": fx.get("lower_rate"),
        "예정가격": fx.get("estimate_price"),
        "사정률": fx.get("price_adjustment_rate"),
        "참여업체수": participant_count,
        "순위": fx["rank"],
        "업체명": fx["company_name"],
        "투찰금액": fx.get("bid_amount"),
        "예가대비투찰률(%)": fx.get("bid_rate_vs_estimate"),
        "기초대비투찰률(%)": fx.get("bid_rate_vs_base"),
        "기초대비사정률(%)": fx.get("adjust_rate"),
        "투찰일시": fx.get("bid_date"),
    }
    return [mapping.get(name) for name in headers]


def merge_master(master_path: Path, rows: list[dict]) -> None:
    wb = load_workbook(master_path)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    notice_col = headers.index("공고번호")

    by_notice: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        by_notice[row["notice_number"]].append(row)

    drop_keys = set()
    for notice in by_notice:
        full, base = notice_keys(notice)
        drop_keys.add(full)
        drop_keys.add(base)

    kept = [headers]
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        number = str(excel_row[notice_col] or "")
        if number in drop_keys:
            continue
        kept.append(list(excel_row))

    ws.delete_rows(1, ws.max_row)
    for r_i, row in enumerate(kept, start=1):
        for c_i, value in enumerate(row, start=1):
            ws.cell(r_i, c_i, value)

    for fixture_rows in by_notice.values():
        for fx in fixture_rows:
            ws.append(master_row(headers, fx, len(fixture_rows)))

    wb.save(master_path)


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--jsonl", required=True)
    parser.add_argument("--sqlite", required=True)
    parser.add_argument("--master", required=True)
    args = parser.parse_args()
    rows = load_jsonl(Path(args.jsonl))
    merge_sqlite(Path(args.sqlite), rows)
    merge_master(Path(args.master), rows)
    print(f"replaced {len(rows)} converted rows into sqlite+master")


if __name__ == "__main__":
    main()
