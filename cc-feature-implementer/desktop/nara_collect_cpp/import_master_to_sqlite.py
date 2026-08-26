"""Load 전체_통합_데이터.xlsx into nara-dashboard.db participants."""
from __future__ import annotations

import sqlite3
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

MASTER = Path(r"A:\c\nara-workspace\master\전체_통합_데이터.xlsx")
DB = Path(r"A:\c\nara-workspace\app\nara-dashboard.db")


def parse_rate(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("%", "").replace(",", "").strip()
    try:
        return float(text)
    except ValueError:
        return None


def parse_long(value) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    digits = str(value).replace(",", "").replace("원", "").strip()
    try:
        return int(float(digits))
    except ValueError:
        return None


def main() -> None:
    wb = load_workbook(MASTER, read_only=True, data_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {name: i for i, name in enumerate(headers) if name}

    def col(name: str):
        i = idx.get(name)
        return i

    con = sqlite3.connect(str(DB))
    cur = con.cursor()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    cur.execute("DELETE FROM participants")
    n = 0
    for excel_row in ws.iter_rows(min_row=2, values_only=True):
        notice = str(excel_row[col("공고번호")] or "").strip()
        if not notice:
            continue
        name = str(excel_row[col("공사명")] or "")
        if "영향진단" in name:
            continue
        rank = parse_long(excel_row[col("순위")]) or 0
        open_date = excel_row[col("개찰일시")]
        open_text = str(open_date).replace("/", "-") if open_date else ""
        cur.execute(
            """
            INSERT INTO participants (
                notice_no, notice_number, notice_name, ordering_agency, region, rank,
                company_name, normalized_company_name, open_date, bid_date,
                bid_amount, bid_rate_vs_estimate, bid_rate_vs_base, adjust_rate,
                estimate_price, lower_limit_price, base_amount, price_adjustment_rate,
                lower_rate, lower_diff, biz_no, ceo_name, remark, is_awarded, updated_at
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
            """,
            (
                notice.split("-")[0],
                notice,
                name,
                str(excel_row[col("발주처")] or ""),
                str(excel_row[col("지역제한")] or ""),
                rank,
                str(excel_row[col("업체명")] or ""),
                str(excel_row[col("업체명")] or ""),
                open_text,
                str(excel_row[col("투찰일시")] or "") or None,
                parse_long(excel_row[col("투찰금액")]),
                parse_rate(excel_row[col("예가대비투찰률(%)")]),
                parse_rate(excel_row[col("기초대비투찰률(%)")]),
                parse_rate(excel_row[col("기초대비사정률(%)")]),
                parse_long(excel_row[col("예정가격")]),
                None,
                parse_long(excel_row[col("기초금액")]),
                parse_rate(excel_row[col("사정률")]),
                parse_rate(excel_row[col("투찰률")]),
                None,
                "",
                "",
                "",
                1 if rank == 1 else 0,
                now,
            ),
        )
        n += 1
    con.commit()
    mn, mx, notices = cur.execute(
        "select min(open_date), max(open_date), count(distinct notice_number) from participants"
    ).fetchone()
    dong = cur.execute(
        """
        select min(open_date), max(open_date), count(distinct notice_number)
        from participants
        where company_name like '%동국%' or normalized_company_name like '%동국%'
        """
    ).fetchone()
    con.close()
    print(f"imported {n} rows notices={notices} span={mn}~{mx}")
    print(f"dongguk notices={dong[2]} span={dong[0]}~{dong[1]}")


if __name__ == "__main__":
    main()
