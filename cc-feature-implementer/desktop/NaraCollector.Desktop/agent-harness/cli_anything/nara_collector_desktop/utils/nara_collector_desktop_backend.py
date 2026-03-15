from __future__ import annotations

import json
import os
import shutil
import sqlite3
import subprocess
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any


DEFAULT_COMPANY = "전체경쟁업체"
DEFAULT_REGION = "전국"
DEFAULT_PERIOD = "3m"
DEFAULT_MIN_RATE = 99.5
DEFAULT_MAX_RATE = 100.0
DEFAULT_PRECISION = "0.01"


@dataclass
class ResolvedPaths:
    software_root: Path
    harness_root: Path
    repo_root: Path
    workspace_root: Path
    app_root: Path
    launch_script: Path
    collect_script: Path
    full_pipeline_script: Path
    master_path: Path
    output_root: Path
    analysis_root: Path
    logs_dir: Path
    collection_log: Path
    db_path: Path
    analytics_project: Path
    desktop_project: Path
    desktop_executable: Path | None

    def to_dict(self) -> dict[str, str | None]:
        payload = asdict(self)
        return {key: (str(value) if value is not None else None) for key, value in payload.items()}


@dataclass
class CommandResult:
    command: list[str]
    cwd: str | None
    exit_code: int
    stdout: str
    stderr: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "command": self.command,
            "cwd": self.cwd,
            "exit_code": self.exit_code,
            "stdout": self.stdout,
            "stderr": self.stderr,
        }


@dataclass
class SessionSnapshot:
    workspace_root: str
    source_file: str
    output_root: str
    region: str = DEFAULT_REGION
    period: str = DEFAULT_PERIOD
    min_rate: float = DEFAULT_MIN_RATE
    max_rate: float = DEFAULT_MAX_RATE
    precision: str = DEFAULT_PRECISION
    company: str = DEFAULT_COMPANY
    last_status_path: str = ""
    last_chart_path: str = ""
    last_launch_target: str = ""

    def normalized(self) -> "SessionSnapshot":
        min_rate = min(self.min_rate, self.max_rate)
        max_rate = max(self.min_rate, self.max_rate)
        precision = "0.001" if str(self.precision).strip() == "0.001" else "0.01"
        return SessionSnapshot(
            workspace_root=self.workspace_root,
            source_file=self.source_file,
            output_root=self.output_root,
            region=self.region or DEFAULT_REGION,
            period=self.period or DEFAULT_PERIOD,
            min_rate=min_rate,
            max_rate=max_rate,
            precision=precision,
            company=self.company or DEFAULT_COMPANY,
            last_status_path=self.last_status_path,
            last_chart_path=self.last_chart_path,
            last_launch_target=self.last_launch_target,
        )

    def to_dict(self) -> dict[str, Any]:
        return asdict(self.normalized())


def discover_repo_root(start: Path) -> Path:
    for candidate in [start, *start.parents]:
        if (candidate / ".agents").exists() and (candidate / "cc-feature-implementer").exists():
            return candidate
    raise FileNotFoundError("Could not discover repository root containing .agents and cc-feature-implementer")


def software_root() -> Path:
    return Path(__file__).resolve().parents[4]


def harness_root() -> Path:
    return software_root() / "agent-harness"


def repo_root() -> Path:
    return discover_repo_root(software_root())


def resolve_workspace_root(explicit: str | None = None) -> Path:
    if explicit:
        return Path(explicit).expanduser().resolve()
    env = os.environ.get("NARA_WORKSPACE_ROOT", "").strip()
    if env:
        return Path(env).expanduser().resolve()
    return (repo_root() / "nara-workspace").resolve()


def resolve_paths(workspace_override: str | None = None) -> ResolvedPaths:
    repo = repo_root()
    workspace = resolve_workspace_root(workspace_override)
    candidates = [
        software_root() / "bin" / "Debug" / "net8.0-windows10.0.19041.0" / "win-x64" / "동국문화재연구원_낙찰분석프로그램.exe",
        software_root() / "bin" / "Debug" / "net8.0-windows10.0.19041.0" / "동국문화재연구원_낙찰분석프로그램.exe",
        software_root() / "bin" / "Release" / "net8.0-windows10.0.19041.0" / "win-x64" / "publish" / "동국문화재연구원_낙찰분석프로그램.exe",
    ]
    desktop_executable = next((path for path in candidates if path.exists()), candidates[0])
    return ResolvedPaths(
        software_root=software_root(),
        harness_root=harness_root(),
        repo_root=repo,
        workspace_root=workspace,
        app_root=workspace / "app",
        launch_script=workspace / "scripts" / "launch_nara_current.ps1",
        collect_script=workspace / "scripts" / "run_api_pipeline.sh",
        full_pipeline_script=workspace / "scripts" / "run_full_pipeline.sh",
        master_path=workspace / "master" / "전체_통합_데이터.xlsx",
        output_root=workspace / "analysis",
        analysis_root=workspace / "analysis",
        logs_dir=workspace / "logs",
        collection_log=workspace / "logs" / "collection_history.jsonl",
        db_path=workspace / "app" / "nara-dashboard.db",
        analytics_project=software_root().parent / "NaraCollector.Analytics.Cli" / "NaraCollector.Analytics.Cli.csproj",
        desktop_project=software_root() / "NaraCollector.Desktop.csproj",
        desktop_executable=desktop_executable,
    )


def default_snapshot(workspace_override: str | None = None) -> SessionSnapshot:
    paths = resolve_paths(workspace_override)
    return SessionSnapshot(
        workspace_root=str(paths.workspace_root),
        source_file=str(paths.master_path),
        output_root=str(paths.analysis_root),
    )


def session_dir() -> Path:
    raw = os.environ.get("CLI_ANYTHING_NARA_COLLECTOR_DESKTOP_SESSION_DIR", "").strip()
    if raw:
        path = Path(raw).expanduser().resolve()
    else:
        path = harness_root() / ".session"
    path.mkdir(parents=True, exist_ok=True)
    return path


def session_file() -> Path:
    return session_dir() / "state.json"


def ensure_workspace_dirs(paths: ResolvedPaths) -> None:
    for path in [
        paths.workspace_root,
        paths.app_root,
        paths.workspace_root / "scripts",
        paths.workspace_root / "master",
        paths.workspace_root / "first-stage",
        paths.workspace_root / "source-data",
        paths.workspace_root / "output",
        paths.analysis_root,
        paths.logs_dir,
        paths.workspace_root / "reports",
    ]:
        path.mkdir(parents=True, exist_ok=True)


def run_command(command: list[str], *, cwd: Path | None = None, dry_run: bool = False, timeout: int | None = None) -> CommandResult:
    if dry_run:
        return CommandResult(command=command, cwd=str(cwd) if cwd else None, exit_code=0, stdout="", stderr="")
    completed = subprocess.run(
        command,
        cwd=str(cwd) if cwd else None,
        text=True,
        capture_output=True,
        timeout=timeout,
        check=False,
    )
    return CommandResult(
        command=command,
        cwd=str(cwd) if cwd else None,
        exit_code=completed.returncode,
        stdout=completed.stdout,
        stderr=completed.stderr,
    )


def build_launch_command(paths: ResolvedPaths) -> list[str]:
    return [
        "powershell.exe",
        "-NoProfile",
        "-ExecutionPolicy",
        "Bypass",
        "-File",
        str(paths.launch_script).replace("/mnt/a/c", "A:\\c").replace("/", "\\"),
    ]


def build_collect_command(
    paths: ResolvedPaths,
    *,
    start_month: str | None,
    end_month: str | None,
    months: int | None,
    years: int | None,
    keywords: str | None,
) -> list[str]:
    command = ["/bin/bash", str(paths.collect_script)]
    if start_month:
        command += ["--start-month", start_month]
    if end_month:
        command += ["--end-month", end_month]
    if months is not None:
        command += ["--months", str(months)]
    if years is not None:
        command += ["--years", str(years)]
    if keywords:
        command += ["--keywords", keywords]
    return command


def build_analyze_command(snapshot: SessionSnapshot, paths: ResolvedPaths) -> list[str]:
    return [
        "dotnet",
        "run",
        "--project",
        str(paths.analytics_project),
        "--",
        "--source-file",
        snapshot.source_file,
        "--output-root",
        snapshot.output_root,
        "--region",
        snapshot.region,
        "--period",
        snapshot.period,
        "--min-rate",
        str(snapshot.min_rate),
        "--max-rate",
        str(snapshot.max_rate),
    ]


def build_export_chart_command(snapshot: SessionSnapshot, paths: ResolvedPaths, output_path: str | None, width: int, height: int) -> list[str]:
    target_output = output_path or str(paths.workspace_root / "output" / "charts" / f"{snapshot.region}_{snapshot.period}_{snapshot.precision}.png")
    executable = paths.desktop_executable or (software_root() / "bin" / "Debug" / "net8.0-windows10.0.19041.0" / "win-x64" / "동국문화재연구원_낙찰분석프로그램.exe")
    return [
        str(executable),
        "--export-chart",
        "--input",
        snapshot.source_file,
        "--output",
        target_output,
        "--precision",
        snapshot.precision,
        "--min-rate",
        str(snapshot.min_rate),
        "--max-rate",
        str(snapshot.max_rate),
        "--width",
        str(width),
        "--height",
        str(height),
        "--company",
        snapshot.company,
    ]


def read_latest_collection_entry(paths: ResolvedPaths) -> dict[str, Any] | None:
    if not paths.collection_log.exists():
        return None
    lines = [line.strip() for line in paths.collection_log.read_text(encoding="utf-8", errors="ignore").splitlines() if line.strip()]
    if not lines:
        return None
    for line in reversed(lines):
        try:
            return json.loads(line)
        except json.JSONDecodeError:
            continue
    return None


def read_master_summary(path: Path) -> dict[str, Any]:
    summary: dict[str, Any] = {"path": str(path), "exists": path.exists(), "rows": 0, "notice_count": 0}
    if not path.exists():
        return summary
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        mapping = {str(value).strip(): idx for idx, value in enumerate(header) if value is not None}
        notice_idx = mapping.get("공고번호")
        notice_numbers: set[str] = set()
        rows = 0
        for row in sheet.iter_rows(min_row=2, values_only=True):
            rows += 1
            if notice_idx is not None and row and len(row) > notice_idx and row[notice_idx] is not None:
                notice_numbers.add(str(row[notice_idx]).strip())
        workbook.close()
        summary["rows"] = rows
        summary["notice_count"] = len(notice_numbers)
    except Exception as exc:  # pragma: no cover - defensive
        summary["error"] = str(exc)
    return summary


def read_db_summary(path: Path) -> dict[str, Any]:
    summary: dict[str, Any] = {"path": str(path), "exists": path.exists(), "participants": 0, "collector_notices": 0}
    if not path.exists():
        return summary
    try:
        con = sqlite3.connect(str(path))
        cur = con.cursor()
        cur.execute("SELECT COUNT(*) FROM participants")
        summary["participants"] = int(cur.fetchone()[0])
        try:
            cur.execute("SELECT COUNT(*) FROM collector_notices")
            summary["collector_notices"] = int(cur.fetchone()[0])
        except sqlite3.Error:
            summary["collector_notices"] = 0
        con.close()
    except sqlite3.Error as exc:  # pragma: no cover - defensive
        summary["error"] = str(exc)
    return summary


def read_analysis_status(paths: ResolvedPaths, region: str) -> dict[str, Any]:
    normalized_region = region or DEFAULT_REGION
    path = paths.analysis_root / f"shadow_{normalized_region}" / "pipeline_status.json"
    payload: dict[str, Any] = {"path": str(path), "exists": path.exists()}
    if not path.exists():
        return payload
    try:
        payload["data"] = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:  # pragma: no cover - defensive
        payload["error"] = str(exc)
    return payload


def collect_status_payload(snapshot: SessionSnapshot) -> dict[str, Any]:
    paths = resolve_paths(snapshot.workspace_root)
    ensure_workspace_dirs(paths)
    return {
        "paths": paths.to_dict(),
        "session": snapshot.to_dict(),
        "master": read_master_summary(Path(snapshot.source_file)),
        "db": read_db_summary(paths.db_path),
        "latest_collection": read_latest_collection_entry(paths),
        "analysis": read_analysis_status(paths, snapshot.region),
    }


def executable_exists(name: str) -> bool:
    return shutil.which(name) is not None
